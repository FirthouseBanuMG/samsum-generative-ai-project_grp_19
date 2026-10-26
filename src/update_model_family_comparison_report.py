
from pathlib import Path
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "outputs" / "model_comparison_50_samples"
DOC_DIR = (
    ROOT / "documents" / "03_Implementation_and_Experiments"
    / "05_Model_Family_Comparison"
)

CSV_PATH = OUT_DIR / "model_family_comparison.csv"
MANIFEST_PATH = OUT_DIR / "sample_manifest.json"
OUTPUT_PATH = DOC_DIR / "model_family_evaluation_50_samples_updated.xlsx"

def main():
    df = pd.read_csv(CSV_PATH, dtype={"sample_id": str})
    manifest = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))

    required = {
        "sample_id", "model_name", "model_family", "input_format",
        "reference_summary", "generated_summary", "generation_seconds",
        "rouge1_f1", "rouge2_f1", "rougeL_f1",
    }
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing columns in evaluation CSV: {sorted(missing)}")

    if len(df) != 150:
        raise ValueError(f"Expected 150 records, found {len(df)}")
    if df["sample_id"].nunique() != 50:
        raise ValueError("Expected 50 unique sample IDs")
    if df["generated_summary"].isna().any():
        raise ValueError("Missing generated summaries found")
    if manifest["sample_count"] != 50:
        raise ValueError("Manifest sample count is not 50")
    if not set(df["sample_id"]).issubset(set(manifest["sample_ids"])):
        raise ValueError("CSV contains IDs not present in the manifest")

    summary = (
        df.groupby(["model_name", "model_family", "input_format"], as_index=False)
        .agg(
            Samples=("sample_id", "nunique"),
            Mean_generation_seconds=("generation_seconds", "mean"),
            Mean_generated_words=("generated_summary",
                                  lambda s: s.fillna("").str.split().str.len().mean()),
            ROUGE_1=("rouge1_f1", "mean"),
            ROUGE_2=("rouge2_f1", "mean"),
            ROUGE_L=("rougeL_f1", "mean"),
        )
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Model Summary"
    ws.append(list(summary.columns))
    for row in summary.itertuples(index=False, name=None):
        ws.append(list(row))

    detail = wb.create_sheet("150 Summary Evaluations")
    detail_cols = [
        "sample_id", "model_name", "model_family", "input_format",
        "dialogue", "reference_summary", "generated_summary",
        "generation_seconds", "max_new_tokens", "do_sample", "num_beams",
        "rouge1_f1", "rouge2_f1", "rougeL_f1",
    ]
    # Dialogue is read from the source test dataset by sample ID.
    test_path = ROOT / "data" / "processed" / "samsum-test-processed.csv"
    test_df = pd.read_csv(test_path, dtype={"id": str})
    dialogue_map = test_df.set_index("id")["dialogue"].to_dict()

    detail_df = df.copy()
    detail_df["dialogue"] = detail_df["sample_id"].map(dialogue_map)
    detail.append(detail_cols)
    for row in detail_df[detail_cols].itertuples(index=False, name=None):
        detail.append(list(row))

    methodology = wb.create_sheet("Methodology & Caveats")
    methodology.append(["Item", "Method / limitation"])
    notes = [
        ("Scope", "50 unique test-split dialogues × 3 models = 150 generated summaries."),
        ("Sampling", f"Random sample with seed {manifest['seed']}; pilot IDs excluded."),
        ("Dataset", manifest["dataset"]),
        ("Input formats", "Qwen: constraint_based; FLAN-T5: constraint_based; DistilBART: dialogue_only."),
        ("Generation", "Greedy decoding: do_sample=False, num_beams=1, max_new_tokens=100."),
        ("Metrics", "ROUGE-1/2/L F1 calculated with rouge_score and stemming."),
        ("Timing", "Average generation_seconds reported by the runner; CPU execution in this recorded run."),
        ("Comparability", "Different input formats mean this compares tested model configurations, not architecture alone."),
        ("Limitations", "ROUGE is not a direct factuality measure. Automated scores do not replace human review."),
        ("Quality review", "Inspect generated summaries for factual support, key-event coverage, and speaker attribution before drawing strong conclusions."),
    ]
    for note in notes:
        methodology.append(list(note))

    manifest_sheet = wb.create_sheet("Sample Manifest")
    manifest_sheet.append(["sample_id"])
    for sample_id in manifest["sample_ids"]:
        manifest_sheet.append([sample_id])

    # Basic, consistent formatting.
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="244062")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            values = [len(str(c.value or "")) for c in list(column)[:100]]
            width = min(max(max(values, default=10) + 2, 12), 42)
            sheet.column_dimensions[letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    DOC_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")
    print(f"Rows: {len(df)}")
    print(summary.to_string(index=False))

if __name__ == "__main__":
    main()
